"""Excel'den etiket verilerini okur.

1. satır her zaman sütun başlığıdır; kullanıcı istediği isimleri yazabilir
("A-K Barkod", "M Barkod", "C Barkod" …). Roller önce bilinen başlık adlarıyla,
sonra adında "barkod / barcode / qr" geçen sütunlarla eşleştirilir; kalanlar
soldan sağa boştaki sütunlara atanır (bir sütun iki role verilmez).

Beklenen roller:
  1. Ana barkod (zorunlu)
  2. İkinci barkod (opsiyonel)
  3. Orta metin (opsiyonel)
  4. Alt metin (opsiyonel)
"""

from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path

from openpyxl import load_workbook

from tspl_builder import LabelData

BARKOD_ALIASES = {"barkod", "serino", "seri no", "seri_no", "barcode", "seri numarası"}
AK_ALIASES = {
    "a-k barkod", "ak barkod", "ak_barkod", "a_k_barkod", "a-kbarkod",
    "2. barkod", "2.barkod", "ikinci barkod", "barkod 2", "barkod2", "barcode 2",
}
ITEM_ALIASES = {"itemno", "item no", "item_no", "item", "orta metin", "metin 1"}
PARCA_ALIASES = {
    "parcaad", "parca ad", "parca_ad", "parçaad", "parça ad", "model",
    "alt metin", "metin 2",
}
# Adında bunlardan biri geçen sütun barkod sayılır (ör. "M Barkod", "C Barcode")
BARCODE_WORDS = ("barkod", "barcode", "qr")

# Excel yüklenmeden önce / başlık boşsa görünen evrensel adlar
DEFAULT_HEADERS = {
    "barkod": "Barkod",
    "ak_barkod": "2. Barkod",
    "item_no": "Orta metin",
    "parca_ad": "Alt metin",
}


@dataclass
class ColumnHeaders:
    """Excel 1. satırından gelen görünen sütun adları."""

    barkod: str = DEFAULT_HEADERS["barkod"]
    ak_barkod: str = DEFAULT_HEADERS["ak_barkod"]
    item_no: str = DEFAULT_HEADERS["item_no"]
    parca_ad: str = DEFAULT_HEADERS["parca_ad"]

    def as_dict(self) -> dict[str, str]:
        return asdict(self)

    @classmethod
    def from_dict(cls, data: dict | None) -> ColumnHeaders:
        data = data or {}
        return cls(
            barkod=str(data.get("barkod") or DEFAULT_HEADERS["barkod"]),
            ak_barkod=str(data.get("ak_barkod") or DEFAULT_HEADERS["ak_barkod"]),
            item_no=str(data.get("item_no") or DEFAULT_HEADERS["item_no"]),
            parca_ad=str(data.get("parca_ad") or DEFAULT_HEADERS["parca_ad"]),
        )


# Arayüzde gösterilen evrensel rol adları (Excel başlığından bağımsız)
ROLE_NAMES = {
    "barkod": "Ana barkod",
    "ak_barkod": "2. barkod",
    "item_no": "Orta metin",
    "parca_ad": "Alt metin",
}


def role_label(role: str, headers: "ColumnHeaders") -> str:
    """'2. barkod' — Excel başlığı farklıysa '2. barkod (M Barkod)'."""
    name = ROLE_NAMES[role]
    header = getattr(headers, role, "").strip()
    if header and header.lower() not in {name.lower(), DEFAULT_HEADERS[role].lower()}:
        return f"{name} ({header})"
    return name


@dataclass
class ExcelLoadResult:
    labels: list[LabelData]
    headers: ColumnHeaders


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _find_alias(lower_map: dict[str, int], aliases: set[str], used: set[int]) -> int | None:
    for name in aliases:
        idx = lower_map.get(name)
        if idx is not None and idx not in used:
            return idx
    return None


def _find_barcode_word(header: list[str], used: set[int]) -> int | None:
    for i, h in enumerate(header):
        if i not in used and any(w in h.lower() for w in BARCODE_WORDS):
            return i
    return None


def _first_free(header: list[str], used: set[int]) -> int | None:
    for i in range(len(header)):
        if i not in used:
            return i
    return None


def _header_at(header: list[str], index: int | None, fallback: str) -> str:
    if index is None or index < 0 or index >= len(header):
        return fallback
    name = header[index].strip()
    return name or fallback


def load_labels_from_excel(
    path: str | Path,
    col_barkod: str = "Barkod",
    col_ak: str = "",
    col_item: str = "",
    col_parca: str = "",
) -> ExcelLoadResult:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Excel bulunamadı: {path}")

    wb = load_workbook(path, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return ExcelLoadResult(labels=[], headers=ColumnHeaders())

    header = [_cell_str(c) for c in rows[0]]
    # 1. satır her zaman başlıktır (kullanıcı özel isim yazabilir)
    lower_map = {h.lower(): i for i, h in enumerate(header) if h}

    def names(saved: str, aliases: set[str]) -> set[str]:
        return ({saved.strip().lower()} if saved and saved.strip() else set()) | aliases

    used: set[int] = set()

    def claim(idx: int | None) -> int | None:
        if idx is not None:
            used.add(idx)
        return idx

    # 1) Bilinen / daha önce kaydedilmiş başlık adları
    i_barkod = claim(_find_alias(lower_map, names(col_barkod, BARKOD_ALIASES), used))
    i_ak = claim(_find_alias(lower_map, names(col_ak, AK_ALIASES), used))
    i_item = claim(_find_alias(lower_map, names(col_item, ITEM_ALIASES), used))
    i_parca = claim(_find_alias(lower_map, names(col_parca, PARCA_ALIASES), used))
    # 2) Adında "barkod" geçen sütunlar (ör. "M Barkod", "C Barkod")
    if i_barkod is None:
        i_barkod = claim(_find_barcode_word(header, used))
    if i_ak is None:
        i_ak = claim(_find_barcode_word(header, used))
    # 3) Kalan roller: soldan sağa boştaki sütunlar
    if i_barkod is None:
        i_barkod = claim(_first_free(header, used))
    if i_ak is None:
        i_ak = claim(_first_free(header, used))
    if i_item is None:
        i_item = claim(_first_free(header, used))
    if i_parca is None:
        i_parca = claim(_first_free(header, used))
    data_rows = rows[1:]

    headers = ColumnHeaders(
        barkod=_header_at(header, i_barkod, DEFAULT_HEADERS["barkod"]),
        ak_barkod=_header_at(header, i_ak, DEFAULT_HEADERS["ak_barkod"]),
        item_no=_header_at(header, i_item, DEFAULT_HEADERS["item_no"]),
        parca_ad=_header_at(header, i_parca, DEFAULT_HEADERS["parca_ad"]),
    )

    labels: list[LabelData] = []
    for row in data_rows:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        barkod = _cell_str(
            row[i_barkod] if i_barkod is not None and i_barkod < len(row) else ""
        )
        if not barkod:
            continue
        ak = _cell_str(row[i_ak] if i_ak is not None and i_ak < len(row) else "")
        item = _cell_str(row[i_item] if i_item is not None and i_item < len(row) else "")
        parca = _cell_str(
            row[i_parca] if i_parca is not None and i_parca < len(row) else ""
        )
        labels.append(
            LabelData(barkod=barkod, ak_barkod=ak, item_no=item, parca_ad=parca)
        )

    return ExcelLoadResult(labels=labels, headers=headers)
