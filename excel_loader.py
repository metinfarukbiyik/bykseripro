"""Excel'den etiket verilerini okur.

1. satır her zaman sütun başlığıdır; kullanıcı istediği isimleri yazabilir.
Sütun rolleri (A→barkod, B→A-K, C→metin1, D→metin2) önce bilinen
başlık adlarıyla eşleştirilir; bulunamazsa sırayla A–D kullanılır.

Beklenen roller:
  1. Ana barkod (zorunlu)
  2. İkinci barkod (opsiyonel)
  3. Orta metin (opsiyonel)
  4. Alt metin (opsiyonel)
"""

from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path

from openpyxl import load_workbook

from tspl_builder import LabelData

BARKOD_ALIASES = {"barkod", "serino", "seri no", "seri_no", "barcode"}
AK_ALIASES = {"a-k barkod", "ak barkod", "ak_barkod", "a_k_barkod", "a-kbarkod"}
ITEM_ALIASES = {"itemno", "item no", "item_no", "item"}
PARCA_ALIASES = {"parcaad", "parca ad", "parca_ad", "parçaad", "parça ad", "model"}

DEFAULT_HEADERS = {
    "barkod": "Barkod",
    "ak_barkod": "A-K Barkod",
    "item_no": "ItemNo",
    "parca_ad": "ParcaAd",
}


@dataclass
class ColumnHeaders:
    """Excel 1. satırından gelen görünen sütun adları."""

    barkod: str = "Barkod"
    ak_barkod: str = "A-K Barkod"
    item_no: str = "ItemNo"
    parca_ad: str = "ParcaAd"

    def as_dict(self) -> dict[str, str]:
        return asdict(self)

    @classmethod
    def from_dict(cls, data: dict | None) -> ColumnHeaders:
        data = data or {}
        return cls(
            barkod=str(data.get("barkod") or DEFAULT_HEADERS["barkod"]),
            ak_barkod=str(data.get("ak_barkod") or DEFAULT_HEADERS["ak_barkod"]),
            item_no=str(data.get("item_no") or DEFAULT_HEADERS["item_no"]),
            parca_ad=str(data.get("parca_ad") or DEFAULT_HEADERS["parca_ad"]),
        )


@dataclass
class ExcelLoadResult:
    labels: list[LabelData]
    headers: ColumnHeaders


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _find_col(lower_map: dict[str, int], aliases: set[str], default: int | None) -> int | None:
    for name in aliases:
        if name in lower_map:
            return lower_map[name]
    return default


def _header_at(header: list[str], index: int | None, fallback: str) -> str:
    if index is None or index < 0 or index >= len(header):
        return fallback
    name = header[index].strip()
    return name or fallback


def load_labels_from_excel(
    path: str | Path,
    col_barkod: str = "Barkod",
    col_ak: str = "A-K Barkod",
    col_item: str = "ItemNo",
    col_parca: str = "ParcaAd",
) -> ExcelLoadResult:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Excel bulunamadı: {path}")

    wb = load_workbook(path, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return ExcelLoadResult(labels=[], headers=ColumnHeaders())

    header = [_cell_str(c) for c in rows[0]]
    # 1. satır her zaman başlıktır (kullanıcı özel isim yazabilir)
    lower_map = {h.lower(): i for i, h in enumerate(header) if h}

    i_barkod = _find_col(lower_map, {col_barkod.lower()} | BARKOD_ALIASES, 0)
    i_ak = _find_col(
        lower_map,
        {col_ak.lower()} | AK_ALIASES,
        1 if len(header) > 1 else None,
    )
    i_item = _find_col(
        lower_map,
        {col_item.lower()} | ITEM_ALIASES,
        2 if len(header) > 2 else None,
    )
    i_parca = _find_col(
        lower_map,
        {col_parca.lower()} | PARCA_ALIASES,
        3 if len(header) > 3 else None,
    )
    data_rows = rows[1:]

    headers = ColumnHeaders(
        barkod=_header_at(header, i_barkod, col_barkod or DEFAULT_HEADERS["barkod"]),
        ak_barkod=_header_at(header, i_ak, col_ak or DEFAULT_HEADERS["ak_barkod"]),
        item_no=_header_at(header, i_item, col_item or DEFAULT_HEADERS["item_no"]),
        parca_ad=_header_at(header, i_parca, col_parca or DEFAULT_HEADERS["parca_ad"]),
    )

    labels: list[LabelData] = []
    for row in data_rows:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        barkod = _cell_str(
            row[i_barkod] if i_barkod is not None and i_barkod < len(row) else ""
        )
        if not barkod:
            continue
        ak = _cell_str(row[i_ak] if i_ak is not None and i_ak < len(row) else "")
        item = _cell_str(row[i_item] if i_item is not None and i_item < len(row) else "")
        parca = _cell_str(
            row[i_parca] if i_parca is not None and i_parca < len(row) else ""
        )
        labels.append(
            LabelData(barkod=barkod, ak_barkod=ak, item_no=item, parca_ad=parca)
        )

    return ExcelLoadResult(labels=labels, headers=headers)
